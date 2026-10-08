"""Read DOCX and PDF documents into one format-independent list of blocks.

A block is a heading (level 1 = chapter, level 2 = article), a paragraph or a
table, in reading order. Chunking works only with blocks and never opens files.
The metadata table at the top of each document is returned separately.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import pdfplumber
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from rag.registry import DOCUMENTS, DocumentInfo

BlockKind = Literal["heading", "paragraph", "table"]

METADATA_KEY = "დოკუმენტის კოდი"
HEADING_NUMBER = re.compile(r"^\d+(\.\d+)*\.?\s")
SENTENCE_END = (".", ":", ";", "!", "?")

DOCX_HEADING_LEVELS = {"Heading 1": 1, "Heading 2": 2}
DOCX_SKIPPED_STYLES = {"Title"}
DOCX_LIST_STYLES = {"List Bullet"}

# Measured in the provided PDFs, which all share one template.
PDF_PAGE_FURNITURE_MAX_SIZE = 8.5  # running header (7.4) and footer (8.0)
PDF_TITLE_MIN_SIZE = 18.0  # 19.0
PDF_CHAPTER_MIN_SIZE = 13.0  # 14.0, bold
PDF_ARTICLE_MIN_SIZE = 11.2  # 11.4, bold; body text is 11.0
PDF_PARAGRAPH_GAP = 19.0  # line pitch is ~15.8 inside a paragraph, ~22.3 between paragraphs


@dataclass
class Block:
    kind: BlockKind
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)
    page: int | None = None

    def to_text(self) -> str:
        if self.kind == "table":
            return "\n".join(" | ".join(row) for row in self.rows)
        return self.text


@dataclass
class LoadedDocument:
    info: DocumentInfo
    metadata: dict[str, str]
    blocks: list[Block]


def load_document(info: DocumentInfo) -> LoadedDocument:
    if info.format == "docx":
        metadata, blocks = _load_docx(info.path)
    elif info.format == "pdf":
        metadata, blocks = _load_pdf(info.path)
    else:
        raise ValueError(f"Unsupported document format: {info.file_name}")
    return LoadedDocument(info=info, metadata=metadata, blocks=_merge_split_tables(blocks))


def load_all() -> list[LoadedDocument]:
    return [load_document(info) for info in DOCUMENTS]


def _clean(text: str | None) -> str:
    return " ".join((text or "").split())


def _metadata_from(rows: list[list[str]]) -> dict[str, str] | None:
    if not any(row and row[0] == METADATA_KEY for row in rows):
        return None
    return {row[0]: row[1] for row in rows if len(row) >= 2}


def _merge_split_tables(blocks: list[Block]) -> list[Block]:
    """Join a table that continues on the next page and repeats its header row."""
    merged: list[Block] = []
    for block in blocks:
        prev = merged[-1] if merged else None
        if (
            block.kind == "table"
            and prev is not None
            and prev.kind == "table"
            and block.rows
            and prev.rows
            and block.rows[0] == prev.rows[0]
        ):
            prev.rows.extend(block.rows[1:])
            continue
        merged.append(block)
    return merged


# --- DOCX ---------------------------------------------------------------


def _load_docx(path: Path) -> tuple[dict[str, str], list[Block]]:
    doc = Document(str(path))
    metadata: dict[str, str] = {}
    blocks: list[Block] = []

    # Walk the body directly: doc.paragraphs and doc.tables lose the order between them.
    for element in doc.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(element, doc)
            text = _clean(paragraph.text)
            style = paragraph.style.name if paragraph.style is not None else ""
            if not text or style in DOCX_SKIPPED_STYLES:
                continue
            if style in DOCX_HEADING_LEVELS:
                blocks.append(Block("heading", text, level=DOCX_HEADING_LEVELS[style]))
            elif style in DOCX_LIST_STYLES:
                blocks.append(Block("paragraph", f"- {text}"))
            else:
                blocks.append(Block("paragraph", text))
        elif tag == "tbl":
            rows = [[_clean(cell.text) for cell in row.cells] for row in Table(element, doc).rows]
            table_metadata = _metadata_from(rows)
            if table_metadata and not metadata:
                metadata = table_metadata
            else:
                blocks.append(Block("table", rows=rows))
    return metadata, blocks


# --- PDF ----------------------------------------------------------------


def _load_pdf(path: Path) -> tuple[dict[str, str], list[Block]]:
    metadata: dict[str, str] = {}
    blocks: list[Block] = []

    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            items: list[tuple[float, Block]] = []

            tables = page.find_tables()
            for table in tables:
                rows = [[_clean(cell) for cell in row] for row in table.extract()]
                table_metadata = _metadata_from(rows)
                if table_metadata and not metadata:
                    metadata = table_metadata
                else:
                    items.append((table.bbox[1], Block("table", rows=rows, page=page_no)))

            boxes = [table.bbox for table in tables]
            text_only = page.filter(lambda obj: not _inside_any(obj, boxes))
            items.extend(_page_text_blocks(text_only, page_no))

            items.sort(key=lambda item: item[0])
            for _, block in items:
                _append_across_pages(blocks, block)
    return metadata, blocks


def _inside_any(obj: dict, boxes: list[tuple[float, float, float, float]]) -> bool:
    return any(
        x0 <= obj["x0"] and obj["x1"] <= x1 and top <= obj["top"] and obj["bottom"] <= bottom
        for x0, top, x1, bottom in boxes
    )


def _classify_line(line: dict) -> tuple[str, int]:
    char = line["chars"][0]
    size = char["size"]
    bold = "Bold" in char["fontname"]
    if size <= PDF_PAGE_FURNITURE_MAX_SIZE or size >= PDF_TITLE_MIN_SIZE:
        return "skip", 0
    if bold and size >= PDF_CHAPTER_MIN_SIZE:
        return "heading", 1
    if bold and size >= PDF_ARTICLE_MIN_SIZE:
        return "heading", 2
    return "paragraph", 0


def _page_text_blocks(page, page_no: int) -> list[tuple[float, Block]]:
    """Group the text lines of one page into headings and paragraphs."""
    items: list[tuple[float, Block]] = []
    prev_top: float | None = None

    for line in page.extract_text_lines():
        kind, level = _classify_line(line)
        if kind == "skip":
            continue
        text = _clean(line["text"])
        last = items[-1][1] if items else None
        gap = line["top"] - prev_top if prev_top is not None else None
        prev_top = line["top"]

        continues_heading = (
            kind == "heading"
            and last is not None
            and last.kind == "heading"
            and last.level == level
            and not HEADING_NUMBER.match(text)
        )
        continues_paragraph = (
            kind == "paragraph"
            and last is not None
            and last.kind == "paragraph"
            and gap is not None
            and gap <= PDF_PARAGRAPH_GAP
        )
        if continues_heading or continues_paragraph:
            last.text = f"{last.text} {text}"
        else:
            items.append((line["top"], Block(kind, text, level=level, page=page_no)))
    return items


def _append_across_pages(blocks: list[Block], block: Block) -> None:
    """Join a paragraph that was cut by a page break in the middle of a sentence."""
    last = blocks[-1] if blocks else None
    if (
        block.kind == "paragraph"
        and last is not None
        and last.kind == "paragraph"
        and last.page != block.page
        and not last.text.endswith(SENTENCE_END)
    ):
        last.text = f"{last.text} {block.text}"
        return
    blocks.append(block)
